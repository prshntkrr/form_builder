"""Turning a spreadsheet somebody has into a dashboard data source.

    .xlsx ──read first sheet──> typed PostgreSQL table ──> Select Data Source

The same one-way load the external database import performs, from a file
instead of a connection: the table is created and filled inside one
transaction, so a failure anywhere leaves the database exactly as it was and
there is never a half-filled table to find later.

Two things the dashboard's own rules decide here:

* The table is named `<name>_tabular`, because `data_source_service` discovers
  sources by that suffix and nothing else would ever appear in the picker.
* Column types are inferred from the values rather than declared, since a
  spreadsheet declares nothing. The inference is deliberately timid — a column
  is only a number if every value in it is one — because a column read as text
  merely limits which widgets can use it, while a column wrongly read as a
  number loses the values that did not fit.
"""
import io
import logging
import re
import warnings
from datetime import date, datetime, time as _time_of_day
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Tuple

from psycopg2 import sql
from psycopg2.extras import execute_values

from app.core.database import table_exists, transaction
from app.modules.dashboards.services.data_source_service import TABULAR_SUFFIX

logger = logging.getLogger(__name__)

_MIDNIGHT = _time_of_day(0, 0)

# 16 MB. Larger than the form-definition workbooks the forms module reads,
# because this one carries data rather than a definition.
MAX_WORKBOOK_BYTES = 16 * 1024 * 1024

# Beyond this the import is refused rather than left to run for minutes and
# then produce a table the browser cannot draw anyway. See the row counts in
# `data_source_service` for why a dashboard over a very large table struggles.
MAX_ROWS = 100_000

# A name this application is willing to create: a PostgreSQL identifier that
# needs no quoting to be safe, quoted anyway when it is used. Deliberately the
# same rule as the external database import applies to its destination; the two
# are kept separate rather than shared so neither module depends on the other.
NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,50}$")

RESERVED = {
    "forms", "form_version", "form_media", "form_export", "form_view",
    "app_user", "app_role", "role_permission", "user_session", "project",
    "project_member", "client_catalog", "client_catalog_value", "data_standard",
    "standard_variable", "standard_variable_option", "unit", "data_dictionary",
    "external_import", "external_connection", "dashboard", "dashboard_version",
}


class ExcelSourceError(Exception):
    """The import cannot start, or cannot finish. The message says why."""


# --------------------------------------------------------------------------- #
# naming
# --------------------------------------------------------------------------- #
def destination_table(name: str) -> str:
    """The table this import will create, or a refusal explaining why not."""
    text = str(name or "").strip().lower()

    # Somebody who types "farmer_data_tabular" means the same table as somebody
    # who types "farmer_data"; the suffix is the dashboard's rule, not theirs.
    if text.endswith(TABULAR_SUFFIX):
        text = text[: -len(TABULAR_SUFFIX)]

    if not NAME.match(text):
        raise ExcelSourceError(
            f"'{name}' is not a usable table name. Use letters, digits and "
            "underscores, starting with a letter."
        )

    if text in RESERVED:
        raise ExcelSourceError(
            f"'{text}' is one of this application's own tables. Choose another name."
        )

    return f"{text}{TABULAR_SUFFIX}"


def _column_name(heading: Any, position: int, taken: set) -> str:
    """A column name from a spreadsheet heading, unique within the sheet."""
    text = re.sub(r"[^a-z0-9]+", "_", str(heading or "").strip().lower()).strip("_")

    if not text or text[0].isdigit():
        # A heading that survives none of that still has a position, and a
        # column called column_4 is better than an import that fails.
        text = f"column_{position + 1}"

    text = text[:58]

    candidate = text
    suffix = 2
    while candidate in taken:
        # Two columns headed the same is common in an exported sheet, and the
        # second one is not a reason to refuse the file.
        candidate = f"{text[:55]}_{suffix}"
        suffix += 1

    taken.add(candidate)
    return candidate


# --------------------------------------------------------------------------- #
# reading
# --------------------------------------------------------------------------- #
def read_sheet(data: bytes) -> Tuple[List[Any], List[Tuple]]:
    """The first sheet's header row and its data rows.

    The first non-empty row is the header. Unlike the forms module's reader,
    which knows the template it is given, this one knows nothing about the file
    and so assumes the least.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise ExcelSourceError(
            "openpyxl is not installed. Add it with: pip install openpyxl"
        ) from exc

    try:
        with warnings.catch_warnings():
            # A sheet may carry validation rules openpyxl cannot model. They do
            # not affect the values, which is all this reads.
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True
            )
    except Exception as exc:
        raise ExcelSourceError(f"That file could not be opened as .xlsx: {exc}") from exc

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")

            if not workbook.sheetnames:
                raise ExcelSourceError("That workbook has no sheets.")

            sheet = workbook[workbook.sheetnames[0]]
            rows = sheet.iter_rows(values_only=True)

            header = None
            for row in rows:
                if any(str(cell).strip() for cell in row if cell is not None):
                    header = row
                    break

            if header is None:
                raise ExcelSourceError("That sheet is empty.")

            # Trailing empty columns are an artefact of how the sheet was saved,
            # not columns anybody meant to have.
            width = max(
                (i + 1 for i, cell in enumerate(header)
                 if cell is not None and str(cell).strip()),
                default=0,
            )

            if not width:
                raise ExcelSourceError("That sheet has no column headings.")

            header = list(header[:width])

            body = []
            for row in rows:
                if not any(cell is not None and str(cell).strip() for cell in row):
                    continue  # a blank row between blocks is not a record

                body.append(tuple(row[:width]) + (None,) * (width - len(row)))

                if len(body) > MAX_ROWS:
                    raise ExcelSourceError(
                        f"That sheet has more than {MAX_ROWS:,} rows. "
                        "Import a smaller extract."
                    )
    finally:
        workbook.close()

    if not body:
        raise ExcelSourceError("That sheet has headings but no data rows.")

    return header, body


# --------------------------------------------------------------------------- #
# typing
# --------------------------------------------------------------------------- #
def _looks_boolean(values: List[Any]) -> bool:
    return all(isinstance(value, bool) for value in values)


def _looks_integer(values: List[Any]) -> bool:
    for value in values:
        if isinstance(value, bool):
            return False
        if isinstance(value, int):
            continue
        if isinstance(value, float) and value.is_integer():
            continue
        if isinstance(value, str):
            try:
                int(value.strip())
            except ValueError:
                return False
            continue
        return False
    return True


def _looks_numeric(values: List[Any]) -> bool:
    for value in values:
        if isinstance(value, bool):
            return False
        if isinstance(value, (int, float, Decimal)):
            continue
        if isinstance(value, str):
            try:
                Decimal(value.strip())
            except (InvalidOperation, ValueError):
                return False
            continue
        return False
    return True


def _looks_temporal(values: List[Any]) -> Tuple[bool, bool]:
    """(every value is a date or datetime, none of them carries a time)."""
    dates_only = True
    for value in values:
        if isinstance(value, datetime):
            # openpyxl reads every date cell as a datetime, so a column of
            # plain dates arrives here with midnight attached. Reading that as
            # DATE keeps a date axis from being labelled 00:00:00 throughout.
            if value.time() != _MIDNIGHT:
                dates_only = False
            continue
        if isinstance(value, date):
            continue
        return False, False
    return True, dates_only


def infer_type(values: List[Any]) -> str:
    """The PostgreSQL type for one column, from the values actually present."""
    present = [
        value for value in values
        if value is not None and not (isinstance(value, str) and not value.strip())
    ]

    if not present:
        # A column that is empty throughout still belongs in the table: the
        # sheet has it, and a later extract may fill it.
        return "TEXT"

    if _looks_boolean(present):
        return "BOOLEAN"

    temporal, dates_only = _looks_temporal(present)
    if temporal:
        return "DATE" if dates_only else "TIMESTAMP"

    if _looks_integer(present):
        return "BIGINT"

    if _looks_numeric(present):
        return "NUMERIC"

    return "TEXT"


def _coerce(value: Any, pg_type: str) -> Any:
    """One cell, as the column's type wants it. Blank is always NULL."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None

    if pg_type == "TEXT":
        return str(value)

    if pg_type == "BOOLEAN":
        return bool(value)

    if pg_type == "BIGINT":
        return int(float(value)) if isinstance(value, (str, float)) else int(value)

    if pg_type == "NUMERIC":
        return Decimal(str(value).strip())

    if pg_type == "DATE" and isinstance(value, datetime):
        return value.date()

    return value


# --------------------------------------------------------------------------- #
# loading
# --------------------------------------------------------------------------- #
def import_workbook(data: bytes, table_name: str, imported_by: str = "") -> Dict[str, Any]:
    """Create the table and fill it, or change nothing at all."""
    destination = destination_table(table_name)
    header, body = read_sheet(data)

    taken = set()
    columns = [
        {"name": _column_name(heading, index, taken), "heading": heading}
        for index, heading in enumerate(header)
    ]

    for index, column in enumerate(columns):
        column["type"] = infer_type([row[index] for row in body])

    with transaction() as cur:
        if table_exists(cur, destination):
            raise ExcelSourceError(
                f"A table called '{destination}' already exists here. Choose "
                "another name — nothing is overwritten."
            )

    create = sql.SQL("CREATE TABLE {} ({})").format(
        sql.Identifier(destination),
        sql.SQL(", ").join(
            sql.SQL("{} {}").format(sql.Identifier(c["name"]), sql.SQL(c["type"]))
            for c in columns
        ),
    )
    insert = sql.SQL("INSERT INTO {} ({}) VALUES %s").format(
        sql.Identifier(destination),
        sql.SQL(", ").join(sql.Identifier(c["name"]) for c in columns),
    )

    try:
        values = [
            tuple(_coerce(cell, columns[index]["type"]) for index, cell in enumerate(row))
            for row in body
        ]
    except (ValueError, InvalidOperation, TypeError) as exc:
        # The inference said every value fit and one did not. Refusing names the
        # file as the problem, which is where the fix is.
        raise ExcelSourceError(
            f"A value in that sheet could not be read: {exc}"
        ) from exc

    try:
        # Create and fill together, so a failure leaves no table behind.
        with transaction() as cur:
            cur.execute(create)
            execute_values(cur, insert, values, page_size=500)
    except ExcelSourceError:
        raise
    except Exception as exc:
        logger.exception("Excel import failed: -> %s", destination)
        raise ExcelSourceError(
            "The import could not be completed and nothing was written. "
            f"({type(exc).__name__})"
        ) from exc

    logger.info(
        "dashboards excel import ok: %s rows=%s cols=%s by %s",
        destination, len(values), len(columns), imported_by or "-",
    )

    return {
        "table_name": destination,
        "rows_loaded": len(values),
        "columns_loaded": len(columns),
        "columns": [
            {"name": c["name"], "heading": str(c["heading"] or ""), "type": c["type"]}
            for c in columns
        ],
    }
